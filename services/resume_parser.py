from io import BytesIO


class ResumeParser:
    SUPPORTED_TYPES = {"pdf", "docx", "txt"}

    @staticmethod
    def extract(uploaded_file) -> tuple[str, str | None]:
        """Extract readable text from a Streamlit upload without exposing parser errors."""
        extension = uploaded_file.name.rsplit(".", 1)[-1].lower() if "." in uploaded_file.name else ""
        if extension not in ResumeParser.SUPPORTED_TYPES:
            return "", "Please upload a PDF, DOCX, or TXT resume."
        try:
            content = uploaded_file.getvalue()
            if extension == "txt":
                text = content.decode("utf-8-sig", errors="replace")
                return ResumeParser._clean_text(text), None
            if extension == "pdf":
                from pypdf import PdfReader
                reader = PdfReader(BytesIO(content))
                pages = []
                for page in reader.pages:
                    try:
                        page_text = page.extract_text(extraction_mode="layout") or ""
                    except (TypeError, ValueError):
                        page_text = page.extract_text() or ""
                    pages.append(page_text)
                text = ResumeParser._clean_text("\n".join(pages))
                return text, None if text.strip() else "No selectable text was found. This may be a scanned image PDF; export it as a text-based PDF, DOCX, or TXT file."
            from docx import Document
            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_rows = []
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        table_rows.append(" | ".join(dict.fromkeys(cells)))
            text = ResumeParser._clean_text("\n".join(paragraphs + table_rows))
            return text, None if text.strip() else "No text was found in this DOCX file."
        except Exception:
            return "", "We couldn't read this document. Try exporting it as a text-based PDF, DOCX, or TXT file."

    @staticmethod
    def _clean_text(text: str) -> str:
        """Preserve lines (for resume sections) while removing parser artifacts."""
        lines, seen = [], set()
        for line in text.replace("\x00", "").replace("\u00a0", " ").replace("\r", "\n").split("\n"):
            clean = " ".join(line.split()).strip()
            if clean and clean not in seen:
                lines.append(clean)
                seen.add(clean)
        return "\n".join(lines)
